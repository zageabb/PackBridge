from __future__ import annotations

import csv
import io
import zipfile
from dataclasses import dataclass
from pathlib import Path

import fitz
from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader

from .local_ocr import LocalOCRError, ocr_pdf_bytes


ALLOWED_EXTENSIONS = {".txt", ".md", ".markdown", ".csv", ".pdf", ".docx", ".xlsx", ".xlsm"}
MAX_EXTRACTED_CHARS = 2_000_000
MAX_PDF_PAGES = 500
MAX_WORKBOOK_CELLS = 150_000
MAX_ARCHIVE_ENTRIES = 5_000
MAX_ARCHIVE_EXPANDED_BYTES = 75 * 1024 * 1024
MAX_ARCHIVE_COMPRESSION_RATIO = 100
CHUNK_CHARS = 7000


class DocumentIngestionError(ValueError):
    pass


@dataclass(frozen=True)
class ExtractedChunk:
    position: int
    locator: str
    text: str


@dataclass(frozen=True)
class ExtractedDocument:
    extension: str
    chunks: list[ExtractedChunk]
    page_count: int | None = None

    @property
    def text(self) -> str:
        return "\n\n".join(f"[{chunk.locator}]\n{chunk.text}" for chunk in self.chunks)


def _clean(value) -> str:
    return "" if value is None else str(value).strip()


def _markdown_escape(value) -> str:
    return _clean(value).replace("\n", " ").replace("|", "\\|")


def _rows_to_markdown(rows: list[list]) -> str:
    clean_rows = [
        [_markdown_escape(cell) for cell in row]
        for row in rows
        if any(_clean(cell) for cell in row)
    ]
    if not clean_rows:
        return ""
    width = max(len(row) for row in clean_rows)
    clean_rows = [row + [""] * (width - len(row)) for row in clean_rows]
    return "\n".join(
        [
            "| " + " | ".join(clean_rows[0]) + " |",
            "| " + " | ".join(["---"] * width) + " |",
            *["| " + " | ".join(row) + " |" for row in clean_rows[1:]],
        ]
    )


def _validate_office_container(payload: bytes) -> None:
    try:
        with zipfile.ZipFile(io.BytesIO(payload)) as archive:
            members = archive.infolist()
            expanded = sum(member.file_size for member in members)
            compressed = sum(member.compress_size for member in members)
    except zipfile.BadZipFile as exc:
        raise DocumentIngestionError("The Office document is not a valid ZIP package.") from exc
    if len(members) > MAX_ARCHIVE_ENTRIES:
        raise DocumentIngestionError("The Office document contains too many package entries.")
    if expanded > MAX_ARCHIVE_EXPANDED_BYTES:
        raise DocumentIngestionError("The expanded Office document exceeds the safety limit.")
    if compressed and expanded / compressed > MAX_ARCHIVE_COMPRESSION_RATIO:
        raise DocumentIngestionError("The Office document compression ratio exceeds the safety limit.")


def _chunks_from_text(locator: str, text: str, start_position: int) -> list[ExtractedChunk]:
    chunks = []
    for offset in range(0, len(text), CHUNK_CHARS):
        value = text[offset : offset + CHUNK_CHARS].strip()
        if not value:
            continue
        part = len(chunks) + 1
        label = locator if len(text) <= CHUNK_CHARS else f"{locator}, part {part}"
        chunks.append(ExtractedChunk(start_position + len(chunks), label, value))
    return chunks


def extract_path(
    path: str | Path,
    *,
    ocr_mode: str = "off",
    ocr_language: str = "eng",
) -> ExtractedDocument:
    path = Path(path)
    suffix = path.suffix.casefold()
    if suffix not in ALLOWED_EXTENSIONS:
        raise DocumentIngestionError(
            "Supported document formats are PDF, DOCX, XLSX/XLSM, CSV, TXT and Markdown."
        )
    payload = path.read_bytes()
    if not payload:
        raise DocumentIngestionError("The uploaded document is empty.")

    if suffix == ".pdf":
        try:
            return _extract_pdf(payload)
        except DocumentIngestionError as exc:
            if ocr_mode.casefold() != "tesseract":
                raise
            try:
                pages = ocr_pdf_bytes(payload, language=ocr_language)
            except LocalOCRError as ocr_exc:
                raise DocumentIngestionError(str(ocr_exc)) from ocr_exc
            chunks: list[ExtractedChunk] = []
            for page_number, text in pages:
                chunks.extend(
                    _chunks_from_text(
                        f"Page {page_number}, OCR",
                        text,
                        len(chunks) + 1,
                    )
                )
            if not chunks:
                raise DocumentIngestionError(
                    "Local OCR did not find readable text in this PDF."
                ) from exc
            return ExtractedDocument(
                ".pdf",
                chunks,
                page_count=max(page for page, _ in pages),
            )
    if suffix in {".docx", ".xlsx", ".xlsm"}:
        _validate_office_container(payload)
    if suffix == ".docx":
        return _extract_docx(payload)
    if suffix in {".xlsx", ".xlsm"}:
        return _extract_xlsx(payload, suffix)
    if suffix == ".csv":
        try:
            text = payload.decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise DocumentIngestionError("CSV uploads must use UTF-8 compatible encoding.") from exc
        rows = list(csv.reader(io.StringIO(text)))
        rendered = _rows_to_markdown(rows)
        if not rendered:
            raise DocumentIngestionError("The CSV contains no readable data.")
        return ExtractedDocument(suffix, _chunks_from_text("CSV", rendered, 1))
    try:
        text = payload.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise DocumentIngestionError("Text/Markdown uploads must use UTF-8 compatible encoding.") from exc
    chunks = _chunks_from_text("Document", text.strip(), 1)
    if not chunks:
        raise DocumentIngestionError("The document contains no readable text.")
    return ExtractedDocument(suffix, chunks)


def _point_in_rect(x: float, y: float, rect) -> bool:
    try:
        x0, y0, x1, y1 = rect
        return x0 <= x <= x1 and y0 <= y <= y1
    except (TypeError, ValueError):
        return False


def _extract_pdf_with_pymupdf(payload: bytes) -> ExtractedDocument:
    document = fitz.open(stream=payload, filetype="pdf")
    try:
        if document.page_count > MAX_PDF_PAGES:
            raise DocumentIngestionError(
                f"PDF exceeds the {MAX_PDF_PAGES}-page extraction limit."
            )

        chunks: list[ExtractedChunk] = []
        total = 0

        for page_number in range(1, document.page_count + 1):
            page = document.load_page(page_number - 1)

            tables = []
            try:
                finder = page.find_tables()
                tables = list(finder.tables)
            except Exception:
                # Table recognition is an enhancement. Normal PDF text extraction
                # remains available if a particular page cannot be analysed.
                tables = []

            table_rects = [tuple(table.bbox) for table in tables if table.bbox]

            outside_blocks = []
            for block in page.get_text("blocks"):
                if len(block) < 5:
                    continue
                x0, y0, x1, y1, raw_text = block[:5]
                text = str(raw_text or "").strip()
                if not text:
                    continue
                centre_x = (float(x0) + float(x1)) / 2
                centre_y = (float(y0) + float(y1)) / 2
                if any(_point_in_rect(centre_x, centre_y, rect) for rect in table_rects):
                    continue
                outside_blocks.append(text)

            page_text = "\n".join(outside_blocks).strip()
            if not page_text and not tables:
                page_text = (page.get_text("text") or "").strip()

            if page_text:
                total += len(page_text)
                if total > MAX_EXTRACTED_CHARS:
                    raise DocumentIngestionError(
                        "Extracted document text exceeds the safety limit."
                    )
                chunks.extend(
                    _chunks_from_text(
                        f"Page {page_number}",
                        page_text,
                        len(chunks) + 1,
                    )
                )

            for table_number, table in enumerate(tables, 1):
                try:
                    rows = table.extract()
                except Exception:
                    rows = []
                rendered = _rows_to_markdown(rows or [])
                if not rendered:
                    continue
                total += len(rendered)
                if total > MAX_EXTRACTED_CHARS:
                    raise DocumentIngestionError(
                        "Extracted document text exceeds the safety limit."
                    )
                chunks.extend(
                    _chunks_from_text(
                        f"Page {page_number}, Table {table_number}",
                        rendered,
                        len(chunks) + 1,
                    )
                )

        if not chunks:
            raise DocumentIngestionError(
                "No readable PDF text was found. A local OCR/vision fallback will be required for this document."
            )
        return ExtractedDocument(".pdf", chunks, page_count=document.page_count)
    finally:
        document.close()


def _extract_pdf_with_pypdf(payload: bytes) -> ExtractedDocument:
    reader = PdfReader(io.BytesIO(payload))
    if len(reader.pages) > MAX_PDF_PAGES:
        raise DocumentIngestionError(
            f"PDF exceeds the {MAX_PDF_PAGES}-page extraction limit."
        )
    chunks: list[ExtractedChunk] = []
    total = 0
    for page_number, page in enumerate(reader.pages, 1):
        page_text = (page.extract_text() or "").strip()
        total += len(page_text)
        if total > MAX_EXTRACTED_CHARS:
            raise DocumentIngestionError("Extracted document text exceeds the safety limit.")
        chunks.extend(
            _chunks_from_text(
                f"Page {page_number}",
                page_text,
                len(chunks) + 1,
            )
        )
    if not chunks:
        raise DocumentIngestionError(
            "No readable PDF text was found. A local OCR/vision fallback will be required for this document."
        )
    return ExtractedDocument(".pdf", chunks, page_count=len(reader.pages))


def _extract_pdf(payload: bytes) -> ExtractedDocument:
    if not payload.startswith(b"%PDF-"):
        raise DocumentIngestionError("The file does not have a valid PDF signature.")
    try:
        return _extract_pdf_with_pymupdf(payload)
    except DocumentIngestionError:
        raise
    except Exception:
        # Conservative fallback retained for PDFs that PyMuPDF cannot parse.
        return _extract_pdf_with_pypdf(payload)


def _extract_docx(payload: bytes) -> ExtractedDocument:
    document = Document(io.BytesIO(payload))
    chunks: list[ExtractedChunk] = []
    heading = ""
    total = 0
    blocks: list[tuple[str, str]] = []
    for paragraph_number, paragraph in enumerate(document.paragraphs, 1):
        value = paragraph.text.strip()
        if not value:
            continue
        total += len(value)
        if total > MAX_EXTRACTED_CHARS:
            raise DocumentIngestionError("Extracted document text exceeds the safety limit.")
        if paragraph.style and paragraph.style.name.casefold().startswith("heading"):
            heading = value
        locator = f"Heading: {heading}" if heading else f"Paragraph {paragraph_number}"
        blocks.append((locator, value))
    for table_number, table in enumerate(document.tables, 1):
        rendered = _rows_to_markdown([[cell.text.strip() for cell in row.cells] for row in table.rows])
        if rendered:
            blocks.append((f"Table {table_number}", rendered))
    for locator, value in blocks:
        chunks.extend(_chunks_from_text(locator, value, len(chunks) + 1))
    if not chunks:
        raise DocumentIngestionError("The DOCX contains no readable text or tables.")
    return ExtractedDocument(".docx", chunks)


def _extract_xlsx(payload: bytes, suffix: str) -> ExtractedDocument:
    workbook = load_workbook(
        io.BytesIO(payload),
        read_only=True,
        data_only=True,
        keep_vba=suffix == ".xlsm",
    )
    chunks: list[ExtractedChunk] = []
    cell_count = 0
    try:
        for worksheet in workbook.worksheets:
            rows = []
            for row in worksheet.iter_rows(values_only=True):
                cell_count += len(row)
                if cell_count > MAX_WORKBOOK_CELLS:
                    raise DocumentIngestionError(
                        f"Workbook exceeds the {MAX_WORKBOOK_CELLS:,}-cell extraction limit."
                    )
                rows.append(["" if value is None else str(value) for value in row])
            rendered = _rows_to_markdown(rows)
            if rendered:
                chunks.extend(
                    _chunks_from_text(
                        f"Sheet: {worksheet.title}",
                        rendered,
                        len(chunks) + 1,
                    )
                )
    finally:
        workbook.close()
    if not chunks:
        raise DocumentIngestionError("The workbook contains no readable values.")
    return ExtractedDocument(suffix, chunks)
